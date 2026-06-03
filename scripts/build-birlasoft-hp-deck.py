#!/usr/bin/env python3
"""Build Birlasoft × HP Field AI deck (.pptx). Use --v2 for manager brief."""

from __future__ import annotations

import argparse
import json
import re
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

try:
    from PIL import Image
except ImportError:
    Image = None  # type: ignore

ROOT = Path(__file__).resolve().parents[1]
SANKAR = Path.home() / "Documents" / "sankar"
OUT_DIR = ROOT / "_bmad-output/capabilities"
COPY_V2 = ROOT / "_bmad-output/planning-artifacts/birlasoft-hp-field-ai-deck-copy-v2.json"

RED = RGBColor(0xE5, 0x2B, 0x2B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x4A, 0x4A, 0x4A)
PAGE_BG = RGBColor(0xF7, 0xF7, 0xF5)
BORDER = RGBColor(0xDD, 0xDD, 0xDD)
HEADER_BG = RGBColor(0xEF, 0xEF, 0xEC)
FONT_SANS = "Helvetica Neue"
FONT_SLAB = "Georgia"
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")

# Sally v3 layout: +2pt type, fill slide (minimal dead space)
FONT_BOOST = 2
MARGIN_X = Inches(0.38)
MARGIN_TOP = Inches(0.58)
MARGIN_BOTTOM = Inches(0.38)
GAP_CALLOUT = Inches(0.26)


def resolve_path(rel: str) -> Path | None:
    if rel.startswith("sankar/"):
        p = SANKAR / rel.removeprefix("sankar/")
    else:
        p = ROOT / rel
    return p if p.is_file() else None


def crop_image(src: Path, mode: str) -> Path:
    if not Image or mode == "none":
        return src
    im = Image.open(src)
    w, h = im.size
    if mode == "top":
        box = (0, 0, w, int(h * 0.55))
    elif mode == "center":
        box = (0, int(h * 0.15), w, int(h * 0.65))
    elif mode == "chat":
        box = (0, int(h * 0.12), w, int(h * 0.75))
    else:
        return src
    out = im.crop(box)
    tmp = Path(tempfile.gettempdir()) / f"birlasoft-hp-{src.stem}-{mode}.png"
    out.save(tmp, format="PNG")
    return tmp


def _font(run, size: int, bold: bool = False, color=INK, name: str = FONT_SANS):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def add_rich_text(paragraph, text: str, size: int = 16, color=INK, prefix: str = ""):
    if prefix:
        r0 = paragraph.add_run()
        r0.text = prefix
        _font(r0, size, color=color)
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run()
            r.text = text[pos : m.start()]
            _font(r, size, color=color)
        rb = paragraph.add_run()
        rb.text = m.group(1)
        _font(rb, size, bold=True, color=INK)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run()
        r.text = text[pos:]
        _font(r, size, color=color)


def fill_slide_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PAGE_BG


def add_red_callout(slide, left, top, width, height, title: str, pointer: bool = False, font_size: int = 24):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = RED
    box.line.fill.background()
    band_h = Inches(0.2)
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, band_h)
    band.fill.solid()
    band.fill.fore_color.rgb = WHITE
    band.line.fill.background()
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.38)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    fs = font_size if len(title) < 16 else max(18, font_size - 2)
    _font(r, fs, bold=True, color=WHITE)
    if pointer:
        tri_w = Inches(0.32)
        tri_h = Inches(0.18)
        tri = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE,
            left + (width - tri_w) // 2,
            top + height,
            tri_w,
            tri_h,
        )
        tri.fill.solid()
        tri.fill.fore_color.rgb = RED
        tri.line.fill.background()
        tri.rotation = 180.0


def add_subtitle(slide, left, top, width, text: str, size: int = 15):
    box = slide.shapes.add_textbox(left, top, width, Inches(0.42))
    p = box.text_frame.paragraphs[0]
    add_rich_text(p, text, size=size, color=MUTED)


def add_bullets_rich(slide, left, top, width, height, items: list[str], size: int = 16):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    spacing = Pt(14) if len(items) <= 3 else Pt(10)
    for i, text in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = spacing
        add_rich_text(p, text, size=size, prefix="■  ")


def add_table(
    slide,
    left,
    top,
    width,
    headers: list[str],
    rows: list[list[str]],
    col_widths: list[float] | None = None,
    bottom=None,
    header_fs: int | None = None,
    body_fs: int | None = None,
):
    nrows = len(rows) + 1
    ncols = len(headers)
    header_fs = header_fs or (13 + FONT_BOOST)
    body_fs = body_fs or ((12 + FONT_BOOST) if ncols > 2 else (13 + FONT_BOOST))
    if bottom is not None and bottom > top:
        height = bottom - top
    else:
        height = Inches(0.48) * nrows
    shape = slide.shapes.add_table(nrows, ncols, left, top, width, height)
    table = shape.table
    row_h = int(height / nrows)
    for ri in range(nrows):
        table.rows[ri].height = row_h

    if col_widths:
        for ci, frac in enumerate(col_widths):
            table.columns[ci].width = int(width * frac)
    else:
        for ci in range(ncols):
            table.columns[ci].width = int(width / ncols)

    for ci, h in enumerate(headers):
        cell = table.cell(0, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RED
        cell.margin_left = cell.margin_right = Inches(0.06)
        cell.margin_top = cell.margin_bottom = Inches(0.04)
        p = cell.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = h
        _font(r, header_fs, bold=True, color=WHITE)

    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri + 1, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 == 0 else HEADER_BG
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            p = cell.text_frame.paragraphs[0]
            p.word_wrap = True
            add_rich_text(p, str(val), size=body_fs)
    return top + height


def add_hero_band(slide, left, top, width, lines: list[str], height=None):
    n = min(len(lines), 3)
    band_h = height or Inches(0.35 * n + 0.35)
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, band_h)
    band.fill.solid()
    band.fill.fore_color.rgb = RED
    band.line.fill.background()
    tf = band.text_frame
    tf.margin_left = Inches(0.18)
    tf.margin_top = Inches(0.1)
    for i, line in enumerate(lines[:3]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        add_rich_text(p, line, size=14 + FONT_BOOST, color=WHITE, prefix="")
    return top + band_h


def add_exhibits(slide, exhibits: list, left, top, max_width, max_h=Inches(3.1)):
    if not exhibits:
        return top
    n = min(len(exhibits), 2)
    gap = Inches(0.12)
    each_w = (max_width - gap * (n - 1)) / n
    x = left
    bottom = top
    for ex in exhibits[:2]:
        p = resolve_path(ex.get("path", ""))
        if not p:
            continue
        label = ex.get("label", "")
        label_h = Inches(0.32) if label else 0
        if label:
            lb = slide.shapes.add_textbox(x, top, each_w, label_h)
            lp = lb.text_frame.paragraphs[0]
            add_rich_text(lp, label, size=11 + FONT_BOOST, color=RED)
        pic_top = top + label_h
        pic = slide.shapes.add_picture(str(crop_image(p, ex.get("crop", "none"))), x, pic_top, width=each_w)
        if pic.height > max_h:
            ratio = max_h / pic.height
            pic.height = int(max_h)
            pic.width = int(pic.width * ratio)
        cap = ex.get("intent", "")
        if cap:
            cb = slide.shapes.add_textbox(x, pic_top + pic.height + Inches(0.03), each_w, Inches(0.25))
            add_rich_text(cb.text_frame.paragraphs[0], cap, size=9 + FONT_BOOST, color=MUTED)
        bottom = max(bottom, pic_top + pic.height + Inches(0.35))
        x += each_w + gap
    return bottom


def panel_geometry(prs: Presentation, narrow_callout: bool = True):
    content_bottom = prs.slide_height - MARGIN_BOTTOM
    callout_top = MARGIN_TOP
    callout_h = content_bottom - callout_top
    callout_w = Inches(2.85) if narrow_callout else Inches(3.45)
    text_left = MARGIN_X + callout_w + GAP_CALLOUT
    text_w = prs.slide_width - text_left - MARGIN_X
    return MARGIN_X, callout_w, callout_top, callout_h, text_left, text_w, content_bottom


def blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def build_cover(slide, slide_data: dict, prs: Presentation):
    fill_slide_bg(slide)
    callout_w = Inches(7.4)
    callout_h = Inches(2.65)
    left = (prs.slide_width - callout_w) // 2
    top = Inches(1.85)
    heading, sub, date_line = "HP Field Service AI", "Birlasoft · Agentic AI practice brief", "May 2026 · Internal"
    for block in slide_data.get("blocks", []):
        if block["type"] == "heading":
            heading = block["text"]
        elif block["type"] == "subheading":
            sub = block["text"]
        elif block["type"] == "paragraph":
            date_line = block["text"]
    add_red_callout(slide, left, top, callout_w, callout_h, heading, pointer=True, font_size=30)
    for yoff, txt, sz, fn in [(0.58, sub, 18, FONT_SANS), (1.02, date_line, 16, FONT_SLAB)]:
        box = slide.shapes.add_textbox(left, top + callout_h + Inches(yoff), callout_w, Inches(0.4))
        p = box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = txt
        _font(r, sz, color=MUTED, name=fn)


def build_purpose(slide, slide_data: dict, prs: Presentation):
    fill_slide_bg(slide)
    margin, cw, ct, ch, tl, tw, bottom = panel_geometry(prs, narrow_callout=True)
    add_red_callout(slide, margin, ct, cw, ch, slide_data.get("callout", "Purpose"), font_size=26)
    y = ct
    vision = slide_data.get("vision_line", "")
    vision_h = Inches(1.0) if vision else 0
    bullet_h = bottom - y - vision_h
    add_bullets_rich(slide, tl, y, tw, bullet_h, slide_data.get("bullets", [])[:3], size=16 + FONT_BOOST)
    y = bottom - vision_h
    if vision:
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, tl, y, tw, vision_h)
        box.fill.solid()
        box.fill.fore_color.rgb = HEADER_BG
        box.line.color.rgb = BORDER
        p = box.text_frame.paragraphs[0]
        p.margin_top = Inches(0.12)
        add_rich_text(p, vision, size=12 + FONT_BOOST, color=MUTED)


def build_context(slide, slide_data: dict, prs: Presentation):
    fill_slide_bg(slide)
    margin, cw, ct, ch, tl, tw, bottom = panel_geometry(prs)
    add_red_callout(slide, margin, ct, cw, ch, slide_data.get("callout", "Context"), font_size=26)
    y = ct
    north = slide_data.get("north_star", "")
    if north:
        y = add_hero_band(slide, tl, y, tw, [north], height=Inches(1.08)) + Inches(0.1)
    obj = slide_data.get("objective", "")
    if obj:
        box = slide.shapes.add_textbox(tl, y, tw, Inches(0.48))
        add_rich_text(box.text_frame.paragraphs[0], obj, size=14 + FONT_BOOST)
        y += Inches(0.5)
    intro = slide_data.get("stack_intro", "")
    if intro:
        ib = slide.shapes.add_textbox(tl, y, tw, Inches(0.62))
        add_rich_text(ib.text_frame.paragraphs[0], intro, size=12 + FONT_BOOST, color=MUTED)
        y += Inches(0.64)
    tbl = slide_data.get("table", {})
    tbl_rows = len(tbl.get("rows", [])) + 1 if tbl else 0
    mig = slide_data.get("migration", {})
    mig_rows = 3 if mig else 0
    fixed_rows = tbl_rows + mig_rows
    if fixed_rows:
        table_block_h = bottom - y
        mig_h = int(table_block_h * (mig_rows / fixed_rows)) if mig_rows else 0
        tbl_h = table_block_h - mig_h
    if mig:
        y = add_table(
            slide,
            tl,
            y,
            tw,
            ["", "Platform shift (field app)"],
            [
                [mig.get("from_label", "From"), mig.get("from", "")],
                [mig.get("to_label", "To"), mig.get("to", "")],
            ],
            col_widths=[0.2, 0.8],
            bottom=y + mig_h,
        ) + Inches(0.06)
    if tbl:
        add_table(
            slide,
            tl,
            y,
            tw,
            tbl.get("headers", []),
            tbl.get("rows", []),
            col_widths=[0.28, 0.72],
            bottom=bottom,
        )


def build_problem(slide, slide_data: dict, prs: Presentation):
    fill_slide_bg(slide)
    margin, cw, ct, ch, tl, tw, bottom = panel_geometry(prs, narrow_callout=True)
    add_red_callout(slide, margin, ct, cw, ch, slide_data.get("callout", "Problem"), font_size=26)
    exhibits = slide_data.get("exhibits", [])
    text_bottom = bottom - (Inches(2.95) if exhibits else 0)
    y = ct
    intro = slide_data.get("pcf_intro", "")
    if intro:
        box = slide.shapes.add_textbox(tl, y, tw, Inches(0.58))
        add_rich_text(box.text_frame.paragraphs[0], intro, size=12 + FONT_BOOST)
        y += Inches(0.6)
    d = slide_data.get("disruption", {})
    if d:
        band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, tl, y, tw, Inches(1.18))
        band.fill.solid()
        band.fill.fore_color.rgb = WHITE
        band.line.color.rgb = RED
        band.line.width = Pt(2)
        p = band.text_frame.paragraphs[0]
        p.margin_left = Inches(0.12)
        p.margin_right = Inches(0.12)
        sentence = (
            f"{d.get('users', 'Users')} faced {d.get('problem', 'a problem')} on "
            f"{d.get('pathway', 'a critical pathway')}, which drove {d.get('disruption', 'major disruption')}."
        )
        add_rich_text(p, sentence, size=13 + FONT_BOOST)
        y += Inches(1.22)
    traps = slide_data.get("migration_trap", slide_data.get("bullets", []))
    add_bullets_rich(slide, tl, y, tw, text_bottom - y, traps, size=13 + FONT_BOOST)
    if exhibits:
        ex_top = text_bottom + Inches(0.08)
        add_exhibits(slide, exhibits, tl, ex_top, tw, max_h=bottom - ex_top)


def build_method(slide, slide_data: dict, prs: Presentation):
    fill_slide_bg(slide)
    margin, cw, ct, ch, tl, tw, bottom = panel_geometry(prs, narrow_callout=True)
    add_red_callout(slide, margin, ct, cw, ch, slide_data.get("callout", "Method"), font_size=26)
    y = ct
    fname = slide_data.get("framework_name", "")
    if fname:
        fb = slide.shapes.add_textbox(tl, y, tw, Inches(0.36))
        fr = fb.text_frame.paragraphs[0].add_run()
        fr.text = fname
        _font(fr, 12 + FONT_BOOST, bold=True, color=RED)
        y += Inches(0.38)
    sub = slide_data.get("subtitle", "")
    if sub:
        add_subtitle(slide, tl, y, tw, sub, size=14 + FONT_BOOST)
        y += Inches(0.4)
    headers = ["#", "Phase", "Ritual · output", "Unlocks"]
    rows = []
    for f in slide_data.get("framework", []):
        ritual = f.get("ritual", "")
        out = f.get("deliverable", f.get("answer", ""))
        cell = f"{ritual} · {out}" if ritual and out else (ritual or out)
        rows.append([f.get("step", ""), f.get("title", ""), cell, f.get("feeds", "")])
    add_table(slide, tl, y, tw, headers, rows, col_widths=[0.06, 0.14, 0.52, 0.28], bottom=bottom)


def build_release_plan(slide, slide_data: dict, prs: Presentation):
    fill_slide_bg(slide)
    margin, cw, ct, ch, tl, tw, bottom = panel_geometry(prs, narrow_callout=True)
    add_red_callout(slide, margin, ct, cw, ch, slide_data.get("callout", "Solution"), font_size=26)
    y = ct
    sub = slide_data.get("subtitle", "")
    if sub:
        add_subtitle(slide, tl, y, tw, sub, size=14 + FONT_BOOST)
        y += Inches(0.42)
    headers = ["Sprint", "What moved the needle", "Sprint rituals", "Signal"]
    rows = []
    for rel in slide_data.get("releases", []):
        rituals = rel.get("rituals", rel.get("engineering", ""))
        needle = rel.get("needle", rel.get("features", ""))
        signal = rel.get("signal", rel.get("tier", ""))
        rows.append([rel.get("sprint", ""), needle, rituals, signal])
    add_table(slide, tl, y, tw, headers, rows, col_widths=[0.08, 0.38, 0.36, 0.18], bottom=bottom)


def build_split_table(slide, slide_data: dict, prs: Presentation):
    fill_slide_bg(slide)
    margin, cw, ct, ch, tl, tw, bottom = panel_geometry(prs)
    add_red_callout(slide, margin, ct, cw, ch, slide_data.get("callout", ""), font_size=26)
    y = ct
    sub = slide_data.get("subtitle", "")
    if sub:
        add_subtitle(slide, tl, y, tw, sub, size=14 + FONT_BOOST)
        y += Inches(0.42)
    bullets = slide_data.get("bullets", [])
    tbl = slide_data.get("table", {})
    bullet_h = Inches(1.15) if bullets else 0
    if bullets:
        add_bullets_rich(slide, tl, y, tw, bullet_h, bullets, size=14 + FONT_BOOST)
        y += bullet_h + Inches(0.08)
    if tbl:
        add_table(
            slide,
            tl,
            y,
            tw,
            tbl.get("headers", []),
            tbl.get("rows", []),
            col_widths=[0.32, 0.68],
            bottom=bottom,
        )


def build_impact(slide, slide_data: dict, prs: Presentation):
    fill_slide_bg(slide)
    margin, cw, ct, ch, tl, tw, bottom = panel_geometry(prs, narrow_callout=True)
    add_red_callout(slide, margin, ct, cw, ch, slide_data.get("callout", "Impact"), font_size=26)
    y = ct
    foot = slide_data.get("footnote", "")
    foot_h = Inches(0.42) if foot else 0
    heroes = slide_data.get("heroes", [])
    hero_h = Inches(1.22) if heroes else 0
    if heroes:
        y = add_hero_band(slide, tl, y, tw, heroes, height=hero_h) + Inches(0.1)
    cols = slide_data.get("categories", [])
    card_h = bottom - y - foot_h
    if cols:
        gap = Inches(0.1)
        each_w = (tw - gap * (len(cols) - 1)) / len(cols)
        x = tl
        hdr_h = Inches(0.44)
        for cat in cols:
            card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, each_w, card_h)
            card.fill.solid()
            card.fill.fore_color.rgb = WHITE
            card.line.color.rgb = BORDER
            hdr = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, each_w, hdr_h)
            hdr.fill.solid()
            hdr.fill.fore_color.rgb = HEADER_BG
            hdr.line.fill.background()
            hp = hdr.text_frame.paragraphs[0]
            hr = hp.add_run()
            hr.text = cat.get("name", "")
            _font(hr, 11 + FONT_BOOST, bold=True, color=RED)
            body = slide.shapes.add_textbox(
                x + Inches(0.1), y + hdr_h + Inches(0.06), each_w - Inches(0.2), card_h - hdr_h - Inches(0.12)
            )
            tf = body.text_frame
            tf.word_wrap = True
            for i, row in enumerate(cat.get("rows", [])):
                label, val = row[0], row[1] if len(row) > 1 else ""
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(10)
                add_rich_text(p, f"{label}: {val}", size=11 + FONT_BOOST)
            x += each_w + gap
    if foot:
        fb = slide.shapes.add_textbox(margin, bottom - foot_h + Inches(0.04), prs.slide_width - 2 * margin, foot_h)
        add_rich_text(fb.text_frame.paragraphs[0], foot, size=9 + FONT_BOOST, color=MUTED)


LAYOUT_BUILDERS = {
    "cover": build_cover,
    "purpose": build_purpose,
    "context": build_context,
    "problem": build_problem,
    "method": build_method,
    "release_plan": build_release_plan,
    "split_table": build_split_table,
    "impact": build_impact,
}


def build_v2(out_path: Path):
    data = json.loads(COPY_V2.read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for slide_data in data["slides"]:
        slide = blank(prs)
        layout = slide_data.get("layout", "split_table")
        builder = LAYOUT_BUILDERS.get(layout, build_split_table)
        builder(slide, slide_data, prs)

    prs.save(out_path)
    print(f"Wrote {out_path} ({len(data['slides'])} slides)")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--v2", action="store_true", help="Build manager brief v2")
    args = parser.parse_args()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if args.v2:
        build_v2(OUT_DIR / "birlasoft-hp-field-ai-v2.pptx")
    else:
        print("Use --v2 for current deck. Example:")
        print("  .venv-deck/bin/python scripts/build-birlasoft-hp-deck.py --v2")


if __name__ == "__main__":
    main()
